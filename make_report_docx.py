from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
FIGURES = ROOT / "figures"
OUT = ROOT / "report.docx"


def set_default_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for section in document.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)


def add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.line_spacing = 1.15


def add_code(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_figure(document: Document, filename: str, caption: str) -> None:
    path = FIGURES / filename
    if not path.exists():
        add_paragraph(document, f"График {filename} не найден.")
        return

    document.add_picture(str(path), width=Cm(15))
    paragraph = document.add_paragraph(caption)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


doc = Document()
set_default_style(doc)

title = doc.add_heading("Сравнение обучения нейронной сети с квадратичной функцией потерь и перекрестной энтропией", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Авторы: ").bold = True
p.add_run("Ляхов Ярослав, Суворов Степан\n")
p.add_run("Университет: ").bold = True
p.add_run("Университет Иннополис")

doc.add_heading("Аннотация", level=1)
add_paragraph(
    doc,
    "В работе рассматривается обучение нейронной сети для бинарной классификации. "
    "Сравниваются две функции потерь: квадратичная функция потерь (MSE) и перекрёстная "
    "энтропия (Binary Cross-Entropy). Главное внимание уделено тому, как выбор функции "
    "потерь влияет на обратное распространение ошибки, градиенты, скорость обучения "
    "и качество классификации.",
)

doc.add_heading("Введение", level=1)
add_paragraph(
    doc,
    "При обучении нейронной сети функция потерь задает численную меру ошибки модели. "
    "Метод обратного распространения ошибки использует производную этой функции. "
    "Так находятся градиенты по параметрам сети. Поэтому разные функции потерь могут "
    "давать разный ход обучения даже при одной архитектуре и одних данных.",
)
add_paragraph(
    doc,
    "Для задач классификации часто используется перекрестная энтропия. "
    "Квадратичную функцию потерь тоже можно применить, если сравнивать "
    "предсказанную вероятность с истинной меткой класса. Цель работы - показать, "
    "почему для классификации чаще выбирают перекрёстную энтропию.",
)

doc.add_heading("Теоретическая часть", level=1)
add_paragraph(
    doc,
    "Пусть сеть выдает логит z. Вероятность класса 1 получается через sigmoid:",
)
add_code(doc, "p = sigmoid(z)")

add_paragraph(doc, "Для перекрестной энтропии в бинарной классификации:")
add_code(doc, "L_BCE = -y log(p) - (1 - y) log(1 - p)")
add_paragraph(doc, "Градиент по логиту:")
add_code(doc, "dL_BCE/dz = p - y")

add_paragraph(doc, "Для квадратичной функции потерь:")
add_code(doc, "L_MSE = (p - y)^2")
add_paragraph(doc, "Градиент по логиту:")
add_code(doc, "dL_MSE/dz = 2(p - y)p(1 - p)")

add_paragraph(
    doc,
    "Главное отличие состоит в множителе p(1 - p). Когда sigmoid насыщается, "
    "вероятность близка к 0 или 1. Тогда этот множитель становится малым. "
    "Поэтому MSE может давать слабый градиент даже при уверенно неверном ответе. "
    "У перекрёстной энтропии эта проблема меньше. При уверенной ошибке градиент остается большим.",
)

doc.add_heading("Эксперимент", level=1)
add_paragraph(
    doc,
    "В эксперименте используется задача бинарной классификации make_moons. Это "
    "нелинейно разделимый набор данных. Поэтому для решения применяется многослойный перцептрон.",
)
add_paragraph(doc, "Условия сравнения:")
add_bullets(
    doc,
    [
        "одна и та же обучающая и тестовая выборки;",
        "одна и та же архитектура сети;",
        "одинаковая инициализация весов;",
        "одинаковый оптимизатор SGD с momentum;",
        "различается только функция потерь.",
    ],
)
add_paragraph(doc, "Сравниваются два варианта обучения:")
add_bullets(
    doc,
    [
        "BCE: перекрестная энтропия с логитами;",
        "MSE: квадратичная функция потерь между sigmoid-вероятностью и меткой класса.",
    ],
)

doc.add_heading("Графики и результаты", level=1)
add_figure(doc, "gradient_bce_vs_mse.png", "Рисунок 1 - Сравнение градиентов BCE и MSE по логиту.")
add_paragraph(
    doc,
    "При уверенной ошибке для объекта класса 1 градиент BCE остается близким к -1, "
    "а градиент MSE стремится к 0. Это значит, что перекрёстная энтропия передает "
    "более сильный сигнал при обратном распространении ошибки.",
)

add_figure(doc, "dataset_make_moons.png", "Рисунок 2 - Обучающая выборка make_moons.")
add_figure(doc, "training_bce_vs_mse.png", "Рисунок 3 - Динамика loss, accuracy, F1 и нормы градиента.")
add_figure(doc, "decision_boundaries_bce_vs_mse.png", "Рисунок 4 - Границы классификации после обучения.")

doc.add_heading("Выводы", level=1)
add_numbered(
    doc,
    [
        "Перекрёстная энтропия хорошо подходит для обучения нейронной сети в задаче классификации. Ее градиент по логиту равен p - y.",
        "Квадратичная функция потерь содержит множитель p(1 - p). Из-за него градиент уменьшается при насыщении sigmoid.",
        "В задаче классификации BCE обычно быстрее исправляет уверенные ошибки модели и дает более простой ход обучения.",
        "MSE можно использовать как метрику качества вероятностных прогнозов. Но как основная функция потерь для классификации она обычно менее удобна, чем перекрестная энтропия.",
    ],
)

doc.add_heading("Использованные материалы", level=1)
add_numbered(
    doc,
    [
        "Juan R. Terven et al. Loss Functions and Metrics in Deep Learning, 2024. См. project_base.pdf.",
        "Документация PyTorch: BCEWithLogitsLoss, MSELoss, автоматическое дифференцирование.",
    ],
)

doc.save(OUT)
print(f"Created {OUT}")
